"""
Загрузка плательщиков из таблицы Excel.

Ожидается лист со строкой заголовков и колонками: ФИО, институт, группа,
дата рождения, б/вн. Порядок колонок неважен — они ищутся по названию,
потому что каждый год таблицу собирают заново и колонки переезжают.

Курс в базе не хранится: первая цифра кода группы — это и есть курс на
момент составления таблицы, из него считается год поступления, а дальше
курс растёт сам каждое 1 сентября (см. backend/domain/academic.py).

Запуск — сначала вхолостую, он ничего не пишет:

    python -m backend.tools.import_payers 2026.xlsx

Убедились, что отчёт выглядит здраво, — тогда с записью:

    python -m backend.tools.import_payers 2026.xlsx --apply

Импорт идёт одной транзакцией: либо все записи, либо ни одной. Повторный
запуск того же файла ничего не задваивает — совпадения по ФИО пропускаются.
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from typing import Dict, List, Optional, Sequence, Tuple

from backend.core.database import SessionLocal
from backend.domain.academic import (
    DURATION_YEARS, EducationLevel, academic_year_start, admission_year_for_course,
    is_graduated,
)
from backend.domain.models import Faculty, Payer, PaymentStatus, SystemUser, UserRole
from backend.infrastructure.repositories import AuditRepository, encrypt_payer

# ---------------------------------------------------------------------------
# Разбор таблицы
# ---------------------------------------------------------------------------

# Заголовки пишут по-разному, поэтому сверяем по набору вариантов, приведя
# к нижнему регистру и убрав пробелы с точками.
COLUMN_ALIASES: Dict[str, Tuple[str, ...]] = {
    "fio": ("фио", "фамилияимяотчество", "фамилия", "студент"),
    "faculty": ("институт", "деректорат", "директорат", "факультет", "подразделение"),
    "group": ("группа", "гр", "номергруппы"),
    "birth": ("датарождения", "др", "датарожд", "деньрождения"),
    "budget": ("бвн", "бюджет", "бвнб", "бюджетвнебюджет", "формаоплаты"),
}

# Буквенная часть кода группы, по которой видно магистратуру. Специалитет
# по буквам не отличить, зато его выдаёт курс: пятикурсника не бывает
# на четырёхлетней программе — см. проверку срока обучения ниже.
MASTER_GROUP_LETTERS = frozenset({"мг", "мгв"})

# Код группы: курс, буквы направления, номер. Номер иногда с буквой на конце
# («2-тд-17с»), поэтому она разрешена и сохраняется как есть.
GROUP_RE = re.compile(
    r"^\s*(\d{1,2})\s*[-–—]\s*([А-Яа-яЁё]{1,6})\s*[-–—]\s*(\d{1,3}[А-Яа-яЁё]?)\s*$"
)

BUDGET_YES = {"б", "б.", "бюджет", "бюдж", "б/б"}
BUDGET_NO = {"вн", "вн.", "внб", "внебюджет", "внебюдж", "платно", "п"}

# Границы правдоподобного возраста. Дата за ними — почти наверняка опечатка
# в таблице (Excel любит подставить текущий год, если ввели только день
# и месяц), и молча записывать её в карточку хуже, чем оставить поле пустым.
MIN_AGE_YEARS = 14
MAX_AGE_YEARS = 80


def _normalize_header(value: object) -> str:
    return re.sub(r"[\s./\\-]", "", str(value or "")).strip().lower()


def _clean(value: object) -> Optional[str]:
    """Значение ячейки строкой без лишних пробелов, пустое — в None."""
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


@dataclass
class Row:
    """Одна строка таблицы, уже разобранная."""
    line: int
    last_name: str
    first_name: str
    middle_name: Optional[str]
    faculty_short: Optional[str]
    group_name: Optional[str]
    course: Optional[int]
    education_level: Optional[str]
    birth: Optional[date]
    is_budget: bool
    warnings: List[str] = field(default_factory=list)

    @property
    def full_name(self) -> str:
        return " ".join(p for p in (self.last_name, self.first_name, self.middle_name) if p)

    @property
    def key(self) -> str:
        """По чему считаем, что это тот же человек."""
        return self.full_name.casefold()


@dataclass
class Skipped:
    line: int
    who: str
    reason: str


def _parse_group(raw: Optional[str]) -> Tuple[Optional[str], Optional[int], Optional[str], Optional[str]]:
    """
    Разобрать код группы: «1-МД-7» → ("1-мд-7", 1, "master"/"bachelor", предупреждение).

    Буквы приводятся к нижнему регистру — так их пишет форма добавления
    плательщика, и без этого «1-МД-7» и «1-мд-7» стали бы разными группами
    в подсказках и в разбивке по группам.
    """
    if not raw:
        return None, None, None, "группа не указана"

    match = GROUP_RE.match(raw)
    if not match:
        return None, None, None, f"код группы не разобран: «{raw}»"

    course_text, letters, number = match.groups()
    course = int(course_text)
    letters = letters.lower()
    group_name = f"{course}-{letters}-{number.lower()}"

    level = EducationLevel.MASTER if letters in MASTER_GROUP_LETTERS else EducationLevel.BACHELOR

    warning = None

    # Курс больше срока обучения — значит уровень угадан неверно: пятикурсник
    # по определению не на четырёхлетней программе. Если бы это оставили как
    # есть, человек ушёл бы в архив прямо в момент загрузки и пропал из списков.
    # Берём ближайший уровень, в который курс укладывается; в системе их три,
    # поэтому пятилетняя очно-заочная программа тоже получит «специалитет» —
    # название приблизительное, зато срок верный.
    if course > DURATION_YEARS[level]:
        fitting = [lvl for lvl, years in DURATION_YEARS.items() if years >= course]
        if fitting:
            level = min(fitting, key=lambda lvl: DURATION_YEARS[lvl])
            warning = (f"курс {course} не укладывается в срок обучения — "
                       f"уровень проставлен как «{level.value}», проверьте")
        else:
            warning = (f"курс {course} больше любого срока обучения — "
                       f"проверьте код группы «{raw}»")

    return group_name, course, level.value, warning


def _parse_birth(value: object, today: date) -> Tuple[Optional[date], Optional[str]]:
    """Дату берём только правдоподобную: остальное лучше оставить пустым."""
    if value is None:
        return None, "дата рождения не указана"

    if hasattr(value, "year") and hasattr(value, "month"):
        parsed = date(value.year, value.month, value.day)
    else:
        text = _clean(value) or ""
        for pattern in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
            try:
                from datetime import datetime
                parsed = datetime.strptime(text, pattern).date()
                break
            except ValueError:
                continue
        else:
            return None, f"дата рождения не разобрана: «{text}»"

    age = today.year - parsed.year - ((today.month, today.day) < (parsed.month, parsed.day))
    if not (MIN_AGE_YEARS <= age <= MAX_AGE_YEARS):
        return None, f"дата рождения неправдоподобна ({parsed.isoformat()}, возраст {age}) — оставлено пустым"
    return parsed, None


def _parse_budget(raw: Optional[str]) -> Tuple[bool, Optional[str]]:
    if not raw:
        return False, "бюджет/внебюджет не указан — записано как внебюджет"
    text = raw.strip().lower().rstrip(".")
    if text in BUDGET_YES:
        return True, None
    if text in BUDGET_NO:
        return False, None
    return False, f"непонятная отметка бюджета «{raw}» — записано как внебюджет"


def read_rows(path: str, today: date) -> Tuple[List[Row], List[Skipped]]:
    """Прочитать лист и разобрать строки. Ошибки не прячем — собираем в список."""
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - зависит от окружения
        sys.exit("Нужен openpyxl: pip install openpyxl")

    sheet = openpyxl.load_workbook(path, data_only=True).active

    header_row = None
    columns: Dict[str, int] = {}
    for row in sheet.iter_rows(min_row=1, max_row=10, values_only=False):
        found: Dict[str, int] = {}
        for cell in row:
            header = _normalize_header(cell.value)
            for field_name, aliases in COLUMN_ALIASES.items():
                if header in aliases and field_name not in found:
                    found[field_name] = cell.column - 1
        if "fio" in found and "group" in found:
            header_row, columns = row[0].row, found
            break

    if header_row is None:
        sys.exit(
            "Не нашёл строку заголовков. Нужны колонки с названиями "
            "«ФИО» и «группа» (регистр и пробелы неважны)."
        )

    missing = [name for name in ("faculty", "birth", "budget") if name not in columns]
    if missing:
        print(f"  внимание: не нашёл колонки {missing} — эти поля останутся пустыми")

    rows: List[Row] = []
    skipped: List[Skipped] = []

    for line, values in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        def cell(name: str, _values=values) -> object:
            index = columns.get(name)
            return _values[index] if index is not None and index < len(_values) else None

        # Номер строки — настоящий номер в таблице, чтобы бухгалтер нашёл
        # её в Excel по отчёту, а не считал от начала данных.
        fio = _clean(cell("fio"))
        meaningful = [cell(name) for name in ("fio", "faculty", "group", "birth", "budget")]

        if not any(v is not None for v in meaningful):
            # Пустая строка. В хвосте таблиц такие часто тянутся с одной лишь
            # сквозной нумерацией в первой колонке — жаловаться на них незачем.
            continue

        if not fio:
            skipped.append(Skipped(line, "—", "нет ФИО, но строка не пустая"))
            continue

        parts = fio.split()
        if len(parts) < 2:
            skipped.append(Skipped(line, fio, "в ФИО меньше двух слов"))
            continue

        group_name, course, level, group_warning = _parse_group(_clean(cell("group")))
        birth, birth_warning = _parse_birth(cell("birth"), today)
        is_budget, budget_warning = _parse_budget(_clean(cell("budget")))

        row = Row(
            line=line,
            last_name=parts[0],
            first_name=parts[1],
            middle_name=" ".join(parts[2:]) or None,
            faculty_short=_clean(cell("faculty")),
            group_name=group_name,
            course=course,
            education_level=level,
            birth=birth,
            is_budget=is_budget,
            warnings=[w for w in (group_warning, birth_warning, budget_warning) if w],
        )
        if not row.faculty_short:
            row.warnings.append("институт не указан")
        rows.append(row)

    return rows, skipped


# ---------------------------------------------------------------------------
# Запись в базу
# ---------------------------------------------------------------------------

def _faculty_index(db) -> Dict[str, Faculty]:
    """Справочник институтов по краткому и полному названию, без учёта регистра."""
    index: Dict[str, Faculty] = {}
    for faculty in db.query(Faculty).all():
        for key in (faculty.short_name, faculty.name):
            if key:
                index[key.strip().casefold()] = faculty
    return index


def _acting_user(db, username: Optional[str]) -> Optional[SystemUser]:
    """От чьего имени пишем запись в журнал изменений."""
    query = db.query(SystemUser)
    if username:
        return query.filter(SystemUser.username == username).first()
    return query.filter(
        SystemUser.role == UserRole.ADMIN, SystemUser.is_active.is_(True)
    ).order_by(SystemUser.id).first()


def _print_report(
    rows: Sequence[Row],
    skipped: Sequence[Skipped],
    duplicates: Sequence[Skipped],
    unknown_faculties: Counter,
    today: date,
) -> None:
    base_year = academic_year_start(today)
    print(f"\nУчебный год: {base_year}-{base_year + 1} (курс 1 → год поступления {base_year})")

    by_faculty = Counter(r.faculty_short or "не указан" for r in rows)
    print(f"\nК загрузке: {len(rows)}")
    for name, count in sorted(by_faculty.items(), key=lambda kv: (-kv[1], kv[0])):
        mark = "  ← нет в справочнике" if name in unknown_faculties else ""
        print(f"  {count:4}  {name}{mark}")

    by_level = Counter(r.education_level or "не определён" for r in rows)
    print("\nУровень:", ", ".join(f"{k}: {v}" for k, v in sorted(by_level.items())))

    by_course = Counter(r.course for r in rows if r.course)
    print("Курс:   ", ", ".join(f"{k}: {v}" for k, v in sorted(by_course.items())))
    print(f"Бюджет:  {sum(1 for r in rows if r.is_budget)} | "
          f"внебюджет: {sum(1 for r in rows if not r.is_budget)}")
    print(f"Без даты рождения: {sum(1 for r in rows if r.birth is None)}")

    flagged = [r for r in rows if r.warnings]
    if flagged:
        print(f"\nЗагрузятся, но требуют внимания — {len(flagged)}:")
        for row in flagged:
            print(f"  строка {row.line:>4}  {row.full_name}")
            for warning in row.warnings:
                print(f"              · {warning}")

    if duplicates:
        print(f"\nУже есть в базе, пропускаются — {len(duplicates)}:")
        for item in duplicates:
            print(f"  строка {item.line:>4}  {item.who}")

    if skipped:
        print(f"\nНе загрузятся — {len(skipped)}:")
        for item in skipped:
            print(f"  строка {item.line:>4}  {item.who} — {item.reason}")

    # Последняя проверка перед записью: кто по этим данным уже выпустился.
    # Такая запись попадёт сразу в архив — из списков, должников и статистики
    # она пропадёт, и заметят это нескоро. Лучше сказать об этом заранее.
    graduating = [
        row for row in rows
        if row.course and is_graduated(
            admission_year_for_course(row.course, today), row.education_level, today
        )
    ]
    if graduating:
        print(f"\nПопадут сразу в архив — {len(graduating)}:")
        for row in graduating:
            print(f"  строка {row.line:>4}  {row.full_name} | группа {row.group_name} "
                  f"| уровень {row.education_level}")
        print("  Проверьте курс и уровень: в архиве их не будет ни в списке, "
              "ни в должниках, ни в статистике.")


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="Загрузка плательщиков из таблицы Excel",
    )
    parser.add_argument("path", help="файл .xlsx")
    parser.add_argument("--apply", action="store_true",
                        help="записать в базу (без флага — только показать отчёт)")
    parser.add_argument("--create-faculties", action="store_true",
                        help="заводить институты, которых нет в справочнике")
    parser.add_argument("--as-user", metavar="ЛОГИН",
                        help="от чьего имени писать в журнал (по умолчанию первый администратор)")
    args = parser.parse_args(argv)

    today = date.today()
    rows, skipped = read_rows(args.path, today)

    db = SessionLocal()
    try:
        faculties = _faculty_index(db)
        unknown_faculties = Counter(
            r.faculty_short for r in rows
            if r.faculty_short and r.faculty_short.casefold() not in faculties
        )

        # Совпадения с тем, что уже в базе, и внутри самого файла.
        existing = {
            " ".join(p for p in parts if p).casefold()
            for parts in db.query(Payer.last_name, Payer.first_name, Payer.middle_name).all()
        }
        duplicates: List[Skipped] = []
        fresh: List[Row] = []
        seen: set = set()
        for row in rows:
            if row.key in existing:
                duplicates.append(Skipped(row.line, row.full_name, "уже есть в базе"))
            elif row.key in seen:
                duplicates.append(Skipped(row.line, row.full_name, "повтор внутри файла"))
            else:
                seen.add(row.key)
                fresh.append(row)

        _print_report(fresh, skipped, duplicates, unknown_faculties, today)

        if unknown_faculties and not args.create_faculties:
            print(f"\nИнститутов нет в справочнике: {', '.join(sorted(unknown_faculties))}.")
            print("Записи загрузятся без института. Чтобы завести их сразу — "
                  "добавьте флаг --create-faculties.")

        if not args.apply:
            print(f"\nРежим просмотра. Ничего не записано. Повторите с --apply, "
                  f"чтобы загрузить {len(fresh)} записей.")
            return 0

        if not fresh:
            print("\nЗагружать нечего.")
            return 0

        if args.create_faculties:
            for short in sorted(unknown_faculties):
                faculty = Faculty(name=short, short_name=short[:20], is_active=True)
                db.add(faculty)
                db.flush()
                faculties[short.casefold()] = faculty
                print(f"  заведён институт: {short}")

        acting = _acting_user(db, args.as_user)
        if args.as_user and acting is None:
            print(f"Пользователь «{args.as_user}» не найден.")
            return 1

        for row in fresh:
            faculty = faculties.get(row.faculty_short.casefold()) if row.faculty_short else None
            payer = Payer(
                last_name=row.last_name,
                first_name=row.first_name,
                middle_name=row.middle_name,
                date_of_birth=row.birth,
                faculty_id=faculty.id if faculty else None,
                group_name=row.group_name,
                admission_year=(admission_year_for_course(row.course, today)
                                if row.course else None),
                education_level=row.education_level,
                is_budget=row.is_budget,
                status=PaymentStatus.UNPAID,
                is_active=True,
                created_by=acting.id if acting else None,
            )
            # Дата рождения в базе зашифрована — объект готовится тем же кодом,
            # что и при добавлении через интерфейс.
            encrypt_payer(payer)
            db.add(payer)

        # Одной транзакцией: половина загруженного файла хуже, чем ничего.
        db.commit()
        print(f"\nЗагружено записей: {len(fresh)}")

        if acting:
            AuditRepository(db).record(
                "import", "payer", None, acting.id,
                f"Загружено из таблицы: {len(fresh)} записей", None,
            )
        else:
            print("  в журнал не записано: не нашёл администратора")
        return 0
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
