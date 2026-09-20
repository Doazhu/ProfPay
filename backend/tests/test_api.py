"""
Сквозные проверки API на живой базе.

Здесь проверяется то, ради чего делалась переработка: вход не перебирается,
правка карточки не портит непереданные поля, список и статистика считаются
в SQL, а выпускники не мешают отчётности.
"""
from datetime import date

import pytest

from backend.core.config import settings
from backend.domain.academic import academic_year_start
from backend.tests.conftest import ADMIN_PASSWORD

API = "/api/v1"


def make_payer(client, faculty_id=None, **overrides):
    payload = {
        "last_name": "Ренёв", "first_name": "Александр", "middle_name": "Дмитриевич",
        "date_of_birth": "2007-09-25",
        "email": "me@doazhu.pro", "phone": "+79001234567",
        "telegram": "@doazhu", "vk": "vk.com/doazhu",
        "group_name": "1-мд-35", "department": "ЦИАТ",
        "admission_year": academic_year_start(), "education_level": "bachelor",
        # Внебюджетник по умолчанию: бюджетнику взнос удерживается из стипендии
        # автоматически, и тесты про оплату проверяли бы не то, что задумано.
        # Кому нужен бюджетник — передаёт is_budget=True явно.
        "is_budget": False, "stipend_amount": "2500.00", "budget_percent": "1",
        "notes": "примечание с 'кавычками'",
        "faculty_id": faculty_id,
    }
    payload.update(overrides)
    response = client.post(f"{API}/payers", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


# ---------------------------------------------------------------------------
# Вход
# ---------------------------------------------------------------------------

def test_login_success(client, admin):
    response = client.post(f"{API}/auth/login",
                           json={"username": "admin", "password": ADMIN_PASSWORD})
    assert response.status_code == 200
    assert "access_token" in client.cookies


def test_login_by_email(client, admin):
    """Бухгалтеру привычнее почта, чем логин."""
    response = client.post(f"{API}/auth/login",
                           json={"username": "admin@profpay.site", "password": ADMIN_PASSWORD})
    assert response.status_code == 200


def test_login_wrong_password_counts_down(client, admin):
    response = client.post(f"{API}/auth/login",
                           json={"username": "admin", "password": "неверный"})
    assert response.status_code == 401
    assert "Осталось попыток: 4" in response.json()["detail"]


def test_account_locks_after_five_attempts(client, admin):
    """Пятая неудача блокирует вход — ровно как просили."""
    for _ in range(settings.MAX_LOGIN_ATTEMPTS - 1):
        assert client.post(f"{API}/auth/login",
                           json={"username": "admin", "password": "неверный"}).status_code == 401

    locked = client.post(f"{API}/auth/login", json={"username": "admin", "password": "неверный"})
    assert locked.status_code == 429

    # Даже верный пароль теперь не пускает.
    after = client.post(f"{API}/auth/login",
                        json={"username": "admin", "password": ADMIN_PASSWORD})
    assert after.status_code == 429


def test_successful_login_resets_counter(client, admin):
    client.post(f"{API}/auth/login", json={"username": "admin", "password": "неверный"})
    client.post(f"{API}/auth/login", json={"username": "admin", "password": ADMIN_PASSWORD})

    again = client.post(f"{API}/auth/login", json={"username": "admin", "password": "неверный"})
    assert "Осталось попыток: 4" in again.json()["detail"]


def test_unknown_user_gives_same_message(client, admin):
    """Ответ не должен подсказывать, есть такой логин или нет."""
    response = client.post(f"{API}/auth/login",
                           json={"username": "несуществующий", "password": "любой"})
    assert response.status_code == 401
    assert response.json()["detail"] == "Неверный логин или пароль"


def test_protected_endpoint_requires_auth(client):
    assert client.get(f"{API}/payers").status_code == 401


# ---------------------------------------------------------------------------
# Плательщики
# ---------------------------------------------------------------------------

def test_create_and_read_payer(auth_client, faculty):
    created = make_payer(auth_client, faculty.id)
    fetched = auth_client.get(f"{API}/payers/{created['id']}").json()

    assert fetched["full_name"] == "Ренёв Александр Дмитриевич"
    assert fetched["email"] == "me@doazhu.pro"
    assert fetched["date_of_birth"] == "2007-09-25"
    assert fetched["notes"] == "примечание с 'кавычками'"   # без двойного экранирования
    assert fetched["course"] == 1
    assert fetched["group_code"] == "1-мд-35"
    assert fetched["is_archived"] is False


def test_partial_update_keeps_untouched_fields(auth_client, faculty):
    """
    Тот самый баг: раньше поля, которых нет в запросе, получали второй слой
    шифрования — телефон становился «gAAAAA…», а дата рождения пропадала.
    """
    created = make_payer(auth_client, faculty.id)

    for _ in range(4):  # несколько сохранений подряд — слои не должны копиться
        response = auth_client.put(f"{API}/payers/{created['id']}",
                                   json={"last_name": "Ренёв-Петров"})
        assert response.status_code == 200

    payer = auth_client.get(f"{API}/payers/{created['id']}").json()
    assert payer["last_name"] == "Ренёв-Петров"
    assert payer["email"] == "me@doazhu.pro"
    assert payer["phone"] == "+79001234567"
    assert payer["department"] == "ЦИАТ"
    assert payer["date_of_birth"] == "2007-09-25"
    assert str(payer["stipend_amount"]) == "2500.00"
    assert not payer["decryption_failed"]


def test_sensitive_fields_encrypted_in_db(auth_client, faculty, db):
    """Контакты в базе — шифротекст, ФИО — открытый текст для поиска в SQL."""
    from sqlalchemy import text
    created = make_payer(auth_client, faculty.id)

    row = db.execute(
        text("SELECT last_name, group_name, email, phone, notes FROM payers WHERE id = :i"),
        {"i": created["id"]},
    ).fetchone()

    assert row.last_name == "Ренёв"        # открыто: по нему ищем и сортируем
    assert row.group_name == "1-мд-35"     # открыто
    assert row.email.startswith("gAAAAA")  # зашифровано
    assert row.phone.startswith("gAAAAA")
    assert row.notes.startswith("gAAAAA")


def test_search_and_pagination(auth_client, faculty):
    for i in range(25):
        make_payer(auth_client, faculty.id,
                   last_name=f"Фамилия{i:02d}", first_name="Имя", email=None,
                   group_name=f"1-мд-{i:02d}")

    page = auth_client.get(f"{API}/payers", params={"page": 1, "per_page": 10}).json()
    assert len(page["items"]) == 10
    assert page["total"] == 25
    assert page["pages"] == 3

    second = auth_client.get(f"{API}/payers", params={"page": 2, "per_page": 10}).json()
    assert {p["id"] for p in page["items"]} & {p["id"] for p in second["items"]} == set()

    found = auth_client.get(f"{API}/payers", params={"search": "Фамилия07"}).json()
    assert found["total"] == 1

    by_group = auth_client.get(f"{API}/payers", params={"search": "мд-03"}).json()
    assert by_group["total"] == 1


def test_sorted_by_surname(auth_client, faculty):
    for surname in ("Яковлев", "Абрамов", "Миронов"):
        make_payer(auth_client, faculty.id, last_name=surname, email=None)

    items = auth_client.get(f"{API}/payers").json()["items"]
    assert [p["last_name"] for p in items] == ["Абрамов", "Миронов", "Яковлев"]


def test_archive_filter(auth_client, faculty):
    base = academic_year_start()
    make_payer(auth_client, faculty.id, last_name="Первокурсник",
               admission_year=base, email=None)
    make_payer(auth_client, faculty.id, last_name="Выпустился",
               admission_year=base - 4, email=None)

    active = auth_client.get(f"{API}/payers").json()
    assert [p["last_name"] for p in active["items"]] == ["Первокурсник"]

    archived = auth_client.get(f"{API}/payers", params={"archive": "archived"}).json()
    assert [p["last_name"] for p in archived["items"]] == ["Выпустился"]
    assert archived["items"][0]["is_archived"] is True

    every = auth_client.get(f"{API}/payers", params={"archive": "all"}).json()
    assert every["total"] == 2


def test_graduates_are_not_debtors(auth_client, faculty):
    base = academic_year_start()
    make_payer(auth_client, faculty.id, last_name="Должник",
               admission_year=base, email=None)
    make_payer(auth_client, faculty.id, last_name="Выпустился",
               admission_year=base - 4, email=None)

    debtors = auth_client.get(f"{API}/debtors").json()
    assert [p["last_name"] for p in debtors["items"]] == ["Должник"]


def test_soft_delete_hides_but_keeps(auth_client, faculty):
    created = make_payer(auth_client, faculty.id)
    assert auth_client.delete(f"{API}/payers/{created['id']}").status_code == 200
    assert auth_client.get(f"{API}/payers").json()["total"] == 0
    # Запись на месте, просто скрыта.
    assert auth_client.get(f"{API}/payers/{created['id']}").status_code == 200


# ---------------------------------------------------------------------------
# Платежи и статус
# ---------------------------------------------------------------------------

def test_partial_payment_sets_partial_status(auth_client, faculty, year_settings):
    """120 ₽ из 240 за год — это частичная оплата, а не полная."""
    payer = make_payer(auth_client, faculty.id)

    auth_client.post(f"{API}/payments", json={
        "payer_id": payer["id"], "amount": "120.00",
        "payment_date": str(date.today()), "academic_year": "2025-2026", "semester": "fall",
    })
    assert auth_client.get(f"{API}/payers/{payer['id']}").json()["status"] == "partial"

    auth_client.post(f"{API}/payments", json={
        "payer_id": payer["id"], "amount": "120.00",
        "payment_date": str(date.today()), "academic_year": "2025-2026", "semester": "spring",
    })
    full = auth_client.get(f"{API}/payers/{payer['id']}").json()
    assert full["status"] == "paid"
    assert str(full["total_paid"]) == "240.00"


def test_deleting_payment_returns_status_to_unpaid(auth_client, faculty, year_settings):
    """Раньше человек оставался «Оплачено» с нулевой суммой."""
    payer = make_payer(auth_client, faculty.id)
    payment = auth_client.post(f"{API}/payments", json={
        "payer_id": payer["id"], "amount": "240.00", "payment_date": str(date.today()),
    }).json()

    assert auth_client.get(f"{API}/payers/{payer['id']}").json()["status"] == "paid"

    auth_client.delete(f"{API}/payments/{payment['id']}")
    after = auth_client.get(f"{API}/payers/{payer['id']}").json()
    assert after["status"] == "unpaid"
    assert str(after["total_paid"]) == "0"


def test_future_payment_rejected(auth_client, faculty):
    from datetime import timedelta
    payer = make_payer(auth_client, faculty.id)
    response = auth_client.post(f"{API}/payments", json={
        "payer_id": payer["id"], "amount": "120.00",
        "payment_date": str(date.today() + timedelta(days=1)),
    })
    assert response.status_code == 422


# ---------------------------------------------------------------------------
# Статистика
# ---------------------------------------------------------------------------

def test_dashboard_excludes_graduates(auth_client, faculty, year_settings):
    base = academic_year_start()
    payer = make_payer(auth_client, faculty.id, last_name="Активный",
                       admission_year=base, email=None)
    make_payer(auth_client, faculty.id, last_name="Выпустился",
               admission_year=base - 4, email=None)
    auth_client.post(f"{API}/payments", json={
        "payer_id": payer["id"], "amount": "240.00", "payment_date": str(date.today()),
    })

    stats = auth_client.get(f"{API}/stats/dashboard").json()
    assert stats["total_payers"] == 1
    assert stats["archived_payers"] == 1
    assert stats["paid_count"] == 1
    assert str(stats["total_paid_amount"]) == "240.00"


def test_faculty_stats(auth_client, faculty):
    make_payer(auth_client, faculty.id, email=None)
    rows = auth_client.get(f"{API}/stats/by-faculty").json()
    row = next(r for r in rows if r["faculty_id"] == faculty.id)
    assert row["total_payers"] == 1
    assert row["faculty_name"] == "ИИТА"


def test_faculty_debtors_match_the_dashboard(auth_client, faculty, year_settings):
    """
    «Должники» в разрезе деректоратов и в сводке наверху — одно и то же число.

    Раньше в разрезе считались только совсем не платившие, а в сводке ещё и
    заплатившие часть. На одном экране под одной подписью стояли разные числа.
    """
    not_paid = make_payer(auth_client, faculty.id, last_name="Неплательщик", email=None)
    partial = make_payer(auth_client, faculty.id, last_name="Частичный", email=None)
    auth_client.post(f"{API}/payers/{partial['id']}/payments", json={
        "amount": "50.00", "payment_date": str(date.today()), "semester": "fall",
    })

    dashboard = auth_client.get(f"{API}/stats/dashboard").json()
    by_faculty = next(r for r in auth_client.get(f"{API}/stats/by-faculty").json()
                      if r["faculty_id"] == faculty.id)
    debtors_page = auth_client.get(f"{API}/debtors").json()

    assert dashboard["total_debtors"] == 2
    assert by_faculty["debtors_count"] == 2
    assert debtors_page["total"] == 2
    assert {not_paid["id"], partial["id"]} == {p["id"] for p in debtors_page["items"]}


# ---------------------------------------------------------------------------
# Пользователи и права
# ---------------------------------------------------------------------------

def test_viewer_cannot_create_payer(client, admin, faculty):
    client.post(f"{API}/auth/login", json={"username": "admin", "password": ADMIN_PASSWORD})
    # Требование второго фактора здесь мешает: проверяются права, а не вход.
    client.put(f"{API}/auth/totp/policy", json={"enabled": False})
    client.post(f"{API}/auth/users", json={
        "username": "viewer", "email": "viewer@profpay.site",
        "password": "ПарольПросмотра2026", "full_name": "Только просмотр", "role": "viewer",
    })
    client.post(f"{API}/auth/logout")

    client.post(f"{API}/auth/login", json={"username": "viewer", "password": "ПарольПросмотра2026"})
    response = client.post(f"{API}/payers", json={"last_name": "Тест", "first_name": "Тест"})
    assert response.status_code == 403


def test_create_user_rejects_duplicates(auth_client, admin):
    payload = {"username": "operator", "email": "op@profpay.site",
               "password": "ПарольОператора2026", "full_name": "Оператор", "role": "operator"}
    assert auth_client.post(f"{API}/auth/users", json=payload).status_code == 201

    assert auth_client.post(f"{API}/auth/users", json=payload).status_code == 400
    assert auth_client.post(f"{API}/auth/users", json={
        **payload, "username": "operator2"
    }).status_code == 400  # почта занята


def test_new_user_can_log_in_immediately(auth_client, admin):
    """
    Раньше второй пользователь при первом входе получал собственный
    мастер-ключ и переставал читать чужие записи.
    """
    auth_client.post(f"{API}/auth/users", json={
        "username": "operator", "email": "op@profpay.site",
        "password": "ПарольОператора2026", "full_name": "Оператор", "role": "operator",
    })
    auth_client.post(f"{API}/auth/logout")

    login = auth_client.post(f"{API}/auth/login",
                             json={"username": "operator", "password": "ПарольОператора2026"})
    assert login.status_code == 200
    assert auth_client.get(f"{API}/payers").status_code == 200


def test_operator_sees_data_created_by_admin(auth_client, admin, faculty):
    """Общий ключ: то, что завёл администратор, читает и оператор."""
    make_payer(auth_client, faculty.id)
    auth_client.post(f"{API}/auth/users", json={
        "username": "operator", "email": "op@profpay.site",
        "password": "ПарольОператора2026", "full_name": "Оператор", "role": "operator",
    })
    auth_client.post(f"{API}/auth/logout")
    auth_client.post(f"{API}/auth/login",
                     json={"username": "operator", "password": "ПарольОператора2026"})

    items = auth_client.get(f"{API}/payers").json()["items"]
    assert items[0]["email"] == "me@doazhu.pro"


def test_cannot_demote_last_admin(auth_client, admin):
    response = auth_client.put(f"{API}/auth/users/{admin.id}", json={"role": "viewer"})
    assert response.status_code == 400
    assert "последний администратор" in response.json()["detail"]


def test_cannot_delete_self(auth_client, admin):
    assert auth_client.delete(f"{API}/auth/users/{admin.id}").status_code == 400


def test_admin_can_unlock_user(auth_client, admin, db):
    auth_client.post(f"{API}/auth/users", json={
        "username": "operator", "email": "op@profpay.site",
        "password": "ПарольОператора2026", "full_name": "Оператор", "role": "operator",
    })
    user_id = next(u["id"] for u in auth_client.get(f"{API}/auth/users").json()
                   if u["username"] == "operator")

    for _ in range(settings.MAX_LOGIN_ATTEMPTS):
        auth_client.post(f"{API}/auth/login", json={"username": "operator", "password": "нет"})

    assert auth_client.post(f"{API}/auth/users/{user_id}/unlock").status_code == 200


# ---------------------------------------------------------------------------
# Заголовки безопасности
# ---------------------------------------------------------------------------

def test_security_headers_present(client):
    headers = client.get(f"{API}/health").headers
    assert "script-src 'self'" in headers["content-security-policy"]
    assert "'unsafe-inline'" not in headers["content-security-policy"].split("style-src")[0]
    assert headers["x-frame-options"] == "DENY"
    assert headers["x-content-type-options"] == "nosniff"
    assert "x-xss-protection" not in headers   # заголовок устарел и сам был вектором


# ---------------------------------------------------------------------------
# Выгрузка в Excel
# ---------------------------------------------------------------------------

def test_export_does_not_leak_formulas(auth_client, faculty):
    """
    Примечание вида «=HYPERLINK(...)» не должно стать формулой у бухгалтера.

    Данные вводит один человек, а выгрузку открывает другой — Excel выполнил бы
    формулу из ячейки на его машине.
    """
    import io

    import openpyxl

    make_payer(auth_client, faculty.id,
               last_name="Формулов", email=None,
               notes='=HYPERLINK("http://example.invalid","жми")')

    response = auth_client.get(f"{API}/payers/export")
    assert response.status_code == 200

    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    values = [c.value for row in sheet.iter_rows(min_row=2) for c in row if c.value]
    notes = next(v for v in values if isinstance(v, str) and "HYPERLINK" in v)

    assert notes.startswith("'=")            # обезврежено апострофом
    assert not any(                          # ни одна ячейка не стала формулой
        cell.data_type == "f" for row in sheet.iter_rows() for cell in row
    )


def test_export_keeps_phone_readable(auth_client, faculty):
    """Телефон с плюсом Excel иначе считает формулой и теряет «+»."""
    import io

    import openpyxl

    make_payer(auth_client, faculty.id, last_name="Телефонов", email=None)
    response = auth_client.get(f"{API}/payers/export")
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active

    phones = [c.value for row in sheet.iter_rows(min_row=2) for c in row
              if isinstance(c.value, str) and "79001234567" in c.value]
    assert phones == ["'+79001234567"]


# ---------------------------------------------------------------------------
# Фильтр «неполные данные»
# ---------------------------------------------------------------------------

def test_missing_fields_lists_the_gaps(auth_client, faculty):
    """Отметка в строке перечисляет ровно то, что не заполнено."""
    full = make_payer(auth_client, faculty.id, last_name="Полный", email=None)
    assert auth_client.get(f"{API}/payers/{full['id']}").json()["missing_fields"] == []

    empty = make_payer(auth_client, None, last_name="Пустой", email=None,
                       group_name=None, admission_year=None, date_of_birth=None)
    assert set(auth_client.get(f"{API}/payers/{empty['id']}").json()["missing_fields"]) == {
        "группа", "год поступления", "деректорат", "дата рождения",
    }


def test_incomplete_filter_returns_only_records_with_gaps(auth_client, faculty):
    make_payer(auth_client, faculty.id, last_name="Полный", email=None)
    without_group = make_payer(auth_client, faculty.id, last_name="Безгруппы", email=None,
                               group_name=None, admission_year=None)
    without_birth = make_payer(auth_client, faculty.id, last_name="Бездаты", email=None,
                               date_of_birth=None)

    page = auth_client.get(f"{API}/payers", params={"incomplete": "true"}).json()
    assert page["total"] == 2
    assert {p["id"] for p in page["items"]} == {without_group["id"], without_birth["id"]}

    # Без флага список прежний — фильтр ничего не прячет сам по себе.
    assert auth_client.get(f"{API}/payers").json()["total"] == 3


def test_incomplete_filter_matches_the_row_marker(auth_client, faculty):
    """
    Условие живёт в двух местах: в SQL для фильтра и в свойстве модели для
    отметки в строке. Разойдутся — фильтр начнёт врать, поэтому сверяем.
    """
    make_payer(auth_client, faculty.id, last_name="Полный", email=None)
    make_payer(auth_client, faculty.id, last_name="Безгруппы", email=None,
               group_name=None, admission_year=None)
    make_payer(auth_client, None, last_name="Бездеректората", email=None)
    make_payer(auth_client, faculty.id, last_name="Бездаты", email=None, date_of_birth=None)

    filtered = auth_client.get(f"{API}/payers", params={"incomplete": "true", "per_page": 100}).json()
    everyone = auth_client.get(f"{API}/payers", params={"per_page": 100}).json()

    by_filter = {p["id"] for p in filtered["items"]}
    by_marker = {p["id"] for p in everyone["items"] if p["missing_fields"]}
    assert by_filter == by_marker


def test_incomplete_filter_combines_with_other_filters(auth_client, faculty):
    make_payer(auth_client, faculty.id, last_name="Иванов", email=None, group_name=None,
               admission_year=None)
    make_payer(auth_client, faculty.id, last_name="Петров", email=None, group_name=None,
               admission_year=None)

    page = auth_client.get(f"{API}/payers",
                           params={"incomplete": "true", "search": "Иванов"}).json()
    assert page["total"] == 1
    assert page["items"][0]["last_name"] == "Иванов"


# ---------------------------------------------------------------------------
# Бюджетники: профком с них денег не собирает
# ---------------------------------------------------------------------------

def test_budget_payer_is_paid_without_any_payment(auth_client, faculty, year_settings):
    """
    Взнос бюджетника удерживают из стипендии помимо профкома. Денег профком
    не получает, поэтому платежа быть не должно — но и должником человек не
    числится, пока состоит в профкоме.
    """
    created = make_payer(auth_client, faculty.id, email=None, is_budget=True)

    assert created["status"] == "paid"
    assert auth_client.get(f"{API}/payers/{created['id']}/payments").json() == []
    # В «Собрано средств» денег не прибавилось: их никто не вносил.
    assert auth_client.get(f"{API}/stats/dashboard").json()["total_paid_amount"] == "0"


def test_budget_payer_is_never_a_debtor(auth_client, faculty, year_settings):
    make_payer(auth_client, faculty.id, email=None, last_name="Бюджетный", is_budget=True)
    make_payer(auth_client, faculty.id, email=None, last_name="Платник", is_budget=False)

    debtors = auth_client.get(f"{API}/debtors").json()
    assert debtors["total"] == 1
    assert debtors["items"][0]["last_name"] == "Платник"


def test_marking_budget_clears_the_debt(auth_client, faculty, year_settings):
    created = make_payer(auth_client, faculty.id, email=None, is_budget=False)
    assert created["status"] == "unpaid"

    updated = auth_client.put(f"{API}/payers/{created['id']}", json={"is_budget": True}).json()
    assert updated["status"] == "paid"
    assert auth_client.get(f"{API}/payers/{created['id']}/payments").json() == []


def test_unmarking_budget_returns_the_debt(auth_client, faculty, year_settings):
    """Платил не он, а стипендия: перестал быть бюджетником — снова должник."""
    created = make_payer(auth_client, faculty.id, email=None, is_budget=True)
    updated = auth_client.put(f"{API}/payers/{created['id']}", json={"is_budget": False}).json()

    assert updated["status"] == "unpaid"
    assert auth_client.get(f"{API}/debtors").json()["total"] == 1


def test_budget_payers_can_be_hidden_from_the_list(auth_client, faculty, year_settings):
    make_payer(auth_client, faculty.id, email=None, last_name="Бюджетный", is_budget=True)
    make_payer(auth_client, faculty.id, email=None, last_name="Платник", is_budget=False)

    assert auth_client.get(f"{API}/payers").json()["total"] == 2

    only_paying = auth_client.get(f"{API}/payers", params={"paying_only": "true"}).json()
    assert only_paying["total"] == 1
    assert only_paying["items"][0]["last_name"] == "Платник"


def test_dashboard_splits_members_into_budget_and_paying(auth_client, faculty, year_settings):
    for index in range(3):
        make_payer(auth_client, faculty.id, email=None, last_name=f"Бюджетный{index}",
                   is_budget=True)
    make_payer(auth_client, faculty.id, email=None, last_name="Платник", is_budget=False)

    stats = auth_client.get(f"{API}/stats/dashboard").json()
    assert stats["total_payers"] == 4        # участники целиком
    assert stats["budget_count"] == 3
    assert stats["paying_count"] == 1
    assert stats["total_debtors"] == 1       # долг только у платника


# ---------------------------------------------------------------------------
# Шаблон бюджетника
# ---------------------------------------------------------------------------

def test_template_is_only_a_template_by_default(auth_client, faculty, year_settings):
    """Без подтверждения шаблон не трогает заведённых."""
    created = make_payer(auth_client, faculty.id, email=None, is_budget=True,
                         stipend_amount="2500.00", budget_percent="1")

    saved = auth_client.put(f"{API}/budget-settings", json={
        "default_stipend_amount": "9999", "default_budget_percent": "2",
    }).json()
    assert saved["applied_to"] == 0

    payer = auth_client.get(f"{API}/payers/{created['id']}").json()
    assert str(payer["stipend_amount"]) == "2500.00"


def test_template_can_be_applied_to_everyone(auth_client, faculty, year_settings):
    budget = make_payer(auth_client, faculty.id, email=None, last_name="Бюджетный",
                        is_budget=True, stipend_amount="2500.00", budget_percent="1")
    regular = make_payer(auth_client, faculty.id, email=None, last_name="Обычный",
                         is_budget=False)

    saved = auth_client.put(f"{API}/budget-settings", json={
        "default_stipend_amount": "9999", "default_budget_percent": "2",
        "apply_to_existing": True,
    }).json()
    assert saved["applied_to"] == 1

    changed = auth_client.get(f"{API}/payers/{budget['id']}").json()
    assert str(changed["stipend_amount"]) == "9999.00"
    assert str(changed["budget_percent"]) == "2.00"

    # Небюджетников это не касается — значения остались прежними.
    untouched = auth_client.get(f"{API}/payers/{regular['id']}").json()
    assert str(untouched["stipend_amount"]) == "2500.00"
    assert str(untouched["budget_percent"]) == "1.00"


def test_impact_counts_who_will_be_overwritten(auth_client, faculty, year_settings):
    """Диалог подтверждения должен честно сказать, скольких перезапишет."""
    make_payer(auth_client, faculty.id, email=None, last_name="Совпадает",
               is_budget=True, stipend_amount="9999", budget_percent="2")
    make_payer(auth_client, faculty.id, email=None, last_name="Отличается",
               is_budget=True, stipend_amount="2500", budget_percent="1")
    make_payer(auth_client, faculty.id, email=None, last_name="Пустой",
               is_budget=True, stipend_amount=None, budget_percent=None)

    impact = auth_client.get(f"{API}/budget-settings/impact",
                             params={"stipend": "9999", "percent": "2"}).json()
    assert impact["budget_payers"] == 3
    assert impact["differing"] == 1
    assert impact["without_values"] == 1


# ---------------------------------------------------------------------------
# Последний курс и уход в архив
# ---------------------------------------------------------------------------

def test_final_year_is_the_last_course_of_the_level(auth_client, faculty, year_settings):
    """
    Бакалавр на четвёртом и магистр на втором выпускаются этим летом,
    первокурсник — нет.
    """
    base = academic_year_start()
    first = make_payer(auth_client, faculty.id, email=None, last_name="Первокурсник",
                       admission_year=base, education_level="bachelor")
    fourth = make_payer(auth_client, faculty.id, email=None, last_name="Четверокурсник",
                        admission_year=base - 3, education_level="bachelor")
    master = make_payer(auth_client, faculty.id, email=None, last_name="Магистр",
                        admission_year=base - 1, education_level="master")

    assert first["is_final_year"] is False
    assert fourth["is_final_year"] is True
    assert master["is_final_year"] is True

    finishing = auth_client.get(f"{API}/payers/finishing").json()
    assert {p["last_name"] for p in finishing} == {"Четверокурсник", "Магистр"}


def test_specialist_finishes_a_year_later(auth_client, faculty, year_settings):
    """Специалитет длится пять лет — на четвёртом курсе ему ещё рано в выпуск."""
    base = academic_year_start()
    fourth = make_payer(auth_client, faculty.id, email=None, last_name="Четвёртый",
                        admission_year=base - 3, education_level="specialist")
    fifth = make_payer(auth_client, faculty.id, email=None, last_name="Пятый",
                       admission_year=base - 4, education_level="specialist")

    assert fourth["is_final_year"] is False
    assert fifth["is_final_year"] is True


def test_continuing_to_master_keeps_the_record_active(auth_client, faculty, year_settings):
    """Поступил в магистратуру — запись остаётся, но начинает отсчёт заново."""
    base = academic_year_start()
    payer = make_payer(auth_client, faculty.id, email=None, admission_year=base - 3,
                       education_level="bachelor", group_name="4-мд-7")
    assert payer["is_final_year"] is True

    updated = auth_client.put(f"{API}/payers/{payer['id']}", json={
        "education_level": "master", "admission_year": base, "group_name": "1-мг-2",
    }).json()

    assert updated["course"] == 1
    assert updated["is_archived"] is False
    assert updated["is_final_year"] is False


def test_leaving_puts_the_record_into_the_archive_at_once(auth_client, faculty, year_settings):
    """
    Отчислился — в архив сразу, не дожидаясь формального конца срока обучения.
    """
    base = academic_year_start()
    payer = make_payer(auth_client, faculty.id, email=None, admission_year=base,
                       education_level="bachelor")
    assert payer["is_archived"] is False

    updated = auth_client.put(f"{API}/payers/{payer['id']}", json={
        "archived_at": str(date.today()),
    }).json()
    assert updated["is_archived"] is True

    assert auth_client.get(f"{API}/payers").json()["total"] == 0
    assert auth_client.get(f"{API}/payers", params={"archive": "archived"}).json()["total"] == 1
    assert auth_client.get(f"{API}/stats/dashboard").json()["total_payers"] == 0


def test_archived_budget_payer_leaves_the_counts(auth_client, faculty, year_settings):
    """Вышел из профкома — перестаёт считаться участником."""
    payer = make_payer(auth_client, faculty.id, email=None, is_budget=True)
    assert auth_client.get(f"{API}/stats/dashboard").json()["budget_count"] == 1

    auth_client.put(f"{API}/payers/{payer['id']}", json={"archived_at": str(date.today())})
    stats = auth_client.get(f"{API}/stats/dashboard").json()
    assert stats["budget_count"] == 0
    assert stats["total_payers"] == 0


def test_returning_from_the_archive(auth_client, faculty, year_settings):
    """Отметку можно снять — например, восстановился."""
    base = academic_year_start()
    payer = make_payer(auth_client, faculty.id, email=None, admission_year=base)
    auth_client.put(f"{API}/payers/{payer['id']}", json={"archived_at": str(date.today())})

    returned = auth_client.put(f"{API}/payers/{payer['id']}", json={"archived_at": None}).json()
    assert returned["is_archived"] is False
    assert auth_client.get(f"{API}/payers").json()["total"] == 1


def test_finishing_count_on_the_dashboard(auth_client, faculty, year_settings):
    base = academic_year_start()
    make_payer(auth_client, faculty.id, email=None, last_name="Выпускник",
               admission_year=base - 3, education_level="bachelor")
    make_payer(auth_client, faculty.id, email=None, last_name="Первокурсник",
               admission_year=base, education_level="bachelor")

    assert auth_client.get(f"{API}/stats/dashboard").json()["finishing_count"] == 1


def test_faculty_breakdown_matches_the_top_summary(auth_client, faculty, year_settings):
    """
    Разрез по деректоратам и сводка наверху должны говорить одно и то же:
    участники делятся на бюджет и платников, должники — только платники.
    """
    for index in range(2):
        make_payer(auth_client, faculty.id, email=None, last_name=f"Бюджетный{index}",
                   is_budget=True)
    make_payer(auth_client, faculty.id, email=None, last_name="Платник", is_budget=False)

    stats = auth_client.get(f"{API}/stats/dashboard").json()
    row = next(r for r in auth_client.get(f"{API}/stats/by-faculty").json()
               if r["faculty_id"] == faculty.id)

    assert row["total_payers"] == stats["total_payers"] == 3
    assert row["budget_count"] == stats["budget_count"] == 2
    assert row["paying_count"] == stats["paying_count"] == 1
    assert row["debtors_count"] == stats["total_debtors"] == 1
    assert row["budget_count"] + row["paying_count"] == row["total_payers"]


def test_startup_marks_existing_budget_payers_as_paid(auth_client, faculty, year_settings, db):
    """
    Записи, заведённые в обход карточки — загруженные из таблицы, например, —
    про правило «бюджетник числится оплаченным» не знают. Приводим их к нему
    на старте приложения: иначе в списке у них висит «Не оплачено», хотя
    должниками они не считаются.
    """
    from backend.core.database import init_db
    from backend.domain.models import Payer, PaymentStatus

    db.add_all([
        Payer(last_name="Бюджетный", first_name="Из", is_budget=True, is_active=True,
              status=PaymentStatus.UNPAID),
        Payer(last_name="Освобождённый", first_name="Из", is_budget=True, is_active=True,
              status=PaymentStatus.EXEMPT),
        Payer(last_name="Платник", first_name="Из", is_budget=False, is_active=True,
              status=PaymentStatus.UNPAID),
    ])
    db.commit()

    init_db()
    db.expire_all()

    by_name = {p.last_name: p.status for p in db.query(Payer).all()}
    assert by_name["Бюджетный"] == PaymentStatus.PAID
    # Освобождение — ручное решение бухгалтера, переписывать его нельзя.
    assert by_name["Освобождённый"] == PaymentStatus.EXEMPT
    assert by_name["Платник"] == PaymentStatus.UNPAID
